"""Shared drawing helpers for the SIH decks."""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
A, ICON = HERE / "assets", HERE / "assets" / "icons"
TEMPLATE = ROOT / "SIH2026-IDEA-Presentation-Format.pptx"
OUT = HERE / "Sentinels_W_PS26142_Idea.pptx"

TEAM, TEAM_ID = "Sentinels_W", "190286"
DEMO = "uchihadzoro.github.io/Deep-Learning-Super-Resolution-Mapping"
CODE = "github.com/UchihaDZoro/Deep-Learning-Super-Resolution-Mapping"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x00, 0x70, 0xC0)      # template footer blue
SAFF = RGBColor(0xE9, 0x7A, 0x1F)      # SIH saffron
GREEN = RGBColor(0x2E, 0x9E, 0x48)     # SIH green
INK = RGBColor(0x1A, 0x1F, 0x2B)
MUTED = RGBColor(0x55, 0x60, 0x70)
TINT = RGBColor(0xF1, 0xF5, 0xFA)
TINT2 = RGBColor(0xFD, 0xF3, 0xEA)
LINE = RGBColor(0xD5, 0xDE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

L, R = 30, 930          # content margins (pt)
TOP, BOTTOM = 92, 490    # below header art, above the blue footer


# ------------------------------------------------------------------ helpers
def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, line=None, margin=0):
    """paras: list of paragraphs; each paragraph is a str or a list of (text, {overrides}) runs.
    A paragraph may also be a dict {"runs": [...], "bullet": True, "size": .., "space_after": ..}."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(margin)
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        spec = p if isinstance(p, dict) else {"runs": p}
        runs = spec["runs"]
        if isinstance(runs, str):
            runs = [(runs, {})]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = spec.get("align", align)
        para.space_after = Pt(spec.get("space_after", space_after))
        if line:
            para.line_spacing = line
        if spec.get("bullet"):
            _bullet(para, spec.get("indent", 11), spec.get("char", "•"), spec.get("bullet_color", BLUE))
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", spec.get("size", size)))
            f.bold = o.get("bold", spec.get("bold", bold))
            f.italic = o.get("italic", False)
            f.underline = o.get("underline", False)
            f.color.rgb = o.get("color", spec.get("color", color))
    return tb


def _bullet(para, indent, char, color):
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent))))
    pPr.set("indent", str(-int(Pt(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    buClr.append(srgb)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buClr); pPr.append(buFont); pPr.append(buChar)


def box(slide, x, y, w, h, fill=TINT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def icon(slide, name, x, y, d, fill=BLUE):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    pad = d * 0.24
    slide.shapes.add_picture(str(ICON / f"{name}.png"), Pt(x + pad), Pt(y + pad), Pt(d - 2 * pad), Pt(d - 2 * pad))
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    p = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w) if w else None, Pt(h) if h else None)
    if border:
        p.line.color.rgb = LINE; p.line.width = Pt(0.75)
    return p


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def pointer(slide, y, label, size=15):
    """Template section pointer, kept verbatim."""
    return text(slide, L, y, R - L, 22, [[("❖ ", {"color": BLUE}), (label, {"underline": True})]],
                size=size, bold=True, color=NAVY)


def team_oval(slide):
    for sh in slide.shapes:
        if sh.name.startswith("Oval") and sh.has_text_frame and "Team" in sh.text_frame.text:
            tf = sh.text_frame
            p0 = tf.paragraphs[0]
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            for r in list(p0.runs)[1:]:
                r._r.getparent().remove(r._r)
            p0.runs[0].text = TEAM
            p0.runs[0].font.size = Pt(11); p0.runs[0].font.bold = True
            p0.runs[0].font.color.rgb = NAVY
            p0.alignment = PP_ALIGN.CENTER
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Pt(2)


def set_title(slide, t, size=None):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = t
            for extra in r[1:]:
                extra._r.getparent().remove(extra._r)
            if size:
                r[0].font.size = Pt(size)
            return sh


def body_box(slide):
    for sh in slide.shapes:
        if sh.shape_type == 17 and sh.name.startswith("TextBox 8"):
            return sh


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


