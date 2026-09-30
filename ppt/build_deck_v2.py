"""In-depth SIH 2026 idea deck (six slides) built on the official template.

python build_deck_v2.py  ->  ppt/Sentinels_W_PS26142_Idea_v2.pptx
Numbers are either measured by our prototype (docs/data/index.json) or published; anything
not yet built is tagged PLANNED / v2.
"""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from deckkit import (A, BLUE, GREEN, INK, LINE, MUTED, NAVY, SAFF, TEMPLATE, TINT, TINT2, WHITE, Presentation,
                     arrow, body_box, box, delete_slide, icon, notes, picture, remove, set_title, team_oval, text)

OUT = A.parent / "Sentinels_W_PS26142_Idea_v2.pptx"
TEAM, TEAM_ID = "Sentinels_W", "190286"
DEMO = "uchihadzoro.github.io/Deep-Learning-Super-Resolution-Mapping"
CODE = "github.com/UchihaDZoro/Deep-Learning-Super-Resolution-Mapping"
TDD = CODE + "/blob/main/TECHNICAL_DESIGN.md"
YOUTUBE = None            # set to the demo-video URL before submission

RED = RGBColor(0xC6, 0x3D, 0x3D)
RED_T = RGBColor(0xFB, 0xEC, 0xEC)
GREEN_T = RGBColor(0xEA, 0xF6, 0xEE)
BLUE_T = RGBColor(0xE8, 0xF1, 0xFA)
SAFF_T = RGBColor(0xFD, 0xF1, 0xE4)
DARK = RGBColor(0x0F, 0x1E, 0x38)
L, R, TOP, BOTTOM = 30, 930, 90, 492


# ------------------------------------------------------------------ extra helpers
def pill(s, x, y, label, fill, color=WHITE, size=6.5, w=None, h=11):
    w = w or (len(label) * size * 0.62 + 10)
    b = box(s, x, y, w, h, fill=fill, radius=0.5)
    text(s, x, y, w, h, [label], size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


BUILT = ("BUILT", GREEN)
PLAN = ("V2", SAFF)


def status(s, x, y, st, right=True):
    lab, col = st
    w = len(lab) * 6.5 * 0.62 + 10
    pill(s, x - w if right else x, y, lab, col, w=w)


def dashed(shape, color):
    shape.line.color.rgb = color
    shape.line.width = Pt(1)
    shape.line.dash_style = MSO_LINE.DASH


def node(s, x, y, w, h, title, sub=None, fill=WHITE, line=LINE, tcol=NAVY, dash=False, tsize=7.8, ssize=6.8, align=PP_ALIGN.CENTER):
    b = box(s, x, y, w, h, fill=fill, line=line, radius=0.12)
    if dash:
        dashed(b, line)
    paras = [[(title, {"bold": True, "color": tcol, "size": tsize})]]
    if sub:
        paras += [[(ln, {"color": MUTED, "size": ssize})] for ln in sub.split("\n")]
    text(s, x + 3, y + 1, w - 6, h - 2, paras, align=align, anchor=MSO_ANCHOR.MIDDLE)
    return b


def tensor(s, x, y, w, h, fill, line, label, sub):
    for k in (2, 1, 0):
        b = box(s, x + 4 * k, y - 4 * k, w - 8, h - 8, fill=fill, line=line, radius=0.06)
    text(s, x, y + 2, w - 8, h - 10, [[(label, {"bold": True, "size": 7.5, "color": NAVY})]] +
         [[(ln, {"size": 6.5, "color": MUTED})] for ln in sub.split("\n")], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def head(s, x, y, w, label, ic, col, size=10.5):
    icon(s, ic, x, y, 16, fill=col)
    text(s, x + 21, y - 1, w - 21, 18, [label], size=size, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)


def ptr(s, x, y, w, label, size=10.5, col=BLUE):
    """Template pointer, kept verbatim."""
    text(s, x, y, w, 16, [[("❖ ", {"color": col}), (label, {"underline": True})]], size=size, bold=True, color=NAVY)


def table(s, x, y, w, colw, rows, header, size=7.3, row_h=20, head_fill=NAVY, zebra=TINT):
    shp = s.shapes.add_table(len(rows) + 1, len(colw), Pt(x), Pt(y), Pt(w), Pt(row_h * (len(rows) + 1)))
    t = shp.table
    tblPr = t._tbl.tblPr
    for el in list(tblPr):
        if el.tag.endswith("tableStyleId"):
            tblPr.remove(el)
    for j, cw in enumerate(colw):
        t.columns[j].width = Pt(cw)
    for i in range(len(rows) + 1):
        t.rows[i].height = Pt(row_h)
        for j in range(len(colw)):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Pt(4); c.margin_top = c.margin_bottom = Pt(1.5)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = header[j] if i == 0 else rows[i - 1][j]
            runs = val if isinstance(val, list) else [(str(val), {})]
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            for r_ in list(p.runs):
                r_._r.getparent().remove(r_._r)
            for tx, o in runs:
                r = p.add_run(); r.text = tx
                r.font.name = "Arial"; r.font.size = Pt(o.get("size", size))
                r.font.bold = True if i == 0 else o.get("bold", False)
                r.font.color.rgb = WHITE if i == 0 else o.get("color", INK)
            c.fill.solid()
            c.fill.fore_color.rgb = head_fill if i == 0 else (zebra if i % 2 == 0 else WHITE)
    return shp


def title(s, t, size=28):
    sh = set_title(s, t, size=size)
    sh.left, sh.top, sh.width, sh.height = Pt(135), Pt(14), Pt(632), Pt(62)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.auto_size = None
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return sh


# ------------------------------------------------------------------ build
prs = Presentation(str(TEMPLATE))
delete_slide(prs, 6)
s1, s2, s3, s4, s5, s6 = prs.slides

# ================================================================== 1 · TITLE
import copy
tb = next(sh for sh in s1.shapes if sh.name == "TextBox 9")
rows = [("Problem Statement ID – ", "26142"),
        ("Problem Statement Title – ", "Deep Learning Based Super Resolution Mapping (SRM) from Medium Resolution Satellite Imageries"),
        ("Theme – ", "Space Technology"), ("PS Category – ", "Software"), ("Team ID – ", TEAM_ID), ("Team Name – ", TEAM)]
tf = tb.text_frame
paras = [p for p in tf.paragraphs if p.text.strip()]
for p in list(tf.paragraphs):
    if not p.text.strip():
        p._p.getparent().remove(p._p)
for p, (lab, val) in zip(paras, rows):
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r0 = runs[0]; r0.text = lab
    r0.font.size = Pt(17); r0.font.bold = True; r0.font.color.rgb = INK
    r1 = copy.deepcopy(r0._r); r0._r.addnext(r1)
    p.runs[1].text = val; p.runs[1].font.bold = False
    p.runs[1].font.color.rgb = NAVY if "Title" in lab else INK
    p.alignment = PP_ALIGN.LEFT; p.line_spacing = 1.1; p.space_before = Pt(0); p.space_after = Pt(14)
tb.top, tb.height = Pt(150), Pt(360)
box(s1, 812, 386, 118, 110, fill=WHITE, line=LINE, radius=0.06)
s1.shapes.add_picture(str(A / "qr_demo.png"), Pt(836), Pt(392), Pt(70), Pt(70))
text(s1, 812, 466, 118, 26, [[("Live prototype", {"bold": True, "color": NAVY})], [("scan to open", {"color": MUTED, "size": 8})]], size=9.5, align=PP_ALIGN.CENTER)
notes(s1, "Team Sentinels_W, PS 26142 (NTRO), Space Technology. The QR code opens the working prototype.")

# ================================================================== 2 · PS & SOLUTION OVERVIEW
team_oval(s2); remove(body_box(s2))
title(s2, "PROBLEM & SOLUTION OVERVIEW", 28)

# A. PS overview + Problem -> Gap -> Solution
box(s2, L, TOP, 290, 80, fill=DARK, radius=0.06)
text(s2, L + 10, TOP + 6, 270, 70, [
    [("PS 26142 · NTRO · SPACE TECHNOLOGY", {"bold": True, "color": RGBColor(0xFF, 0xB4, 0x54), "size": 7.5})],
    {"runs": [("Deep-learning super-resolution mapping from medium-resolution satellite imagery", {"bold": True, "color": WHITE, "size": 9.3})], "space_after": 3},
    [("Asks for: ", {"bold": True, "color": RGBColor(0xC9, 0xD6, 0xEA), "size": 7.4}),
     ("Sentinel-2 10 m → < 4 m · geo & spectral consistency · explicit uncertainty · validation against HR reference", {"color": RGBColor(0xC9, 0xD6, 0xEA), "size": 7.4})],
])
pgs = [("warn", SAFF, SAFF_T, "PROBLEM", "10 m pixels hide small buildings, narrow roads, field bunds and localised damage, which limits classification, change detection and response."),
       ("crosshairs", RED, RED_T, "GAP", "Commercial VHR is costly and must be tasked. Generic SR hallucinates, alters radiometry and reports no confidence."),
       ("check", GREEN, GREEN_T, "SOLUTION", "TRINETRA-SR: physics-constrained 4× SR on all 10 bands, a per-pixel trust map, and it runs anywhere, even offline.")]
gx, gw = L + 300, (R - L - 300 - 2 * 18) / 3
for i, (ic, col, tint, lab, desc) in enumerate(pgs):
    x = gx + i * (gw + 18)
    box(s2, x, TOP, gw, 80, fill=tint, radius=0.08)
    icon(s2, ic, x + 8, TOP + 8, 18, fill=col)
    text(s2, x + 31, TOP + 8, gw - 36, 18, [lab], size=9.5, bold=True, color=col, anchor=MSO_ANCHOR.MIDDLE)
    text(s2, x + 8, TOP + 30, gw - 14, 48, [desc], size=7.6, color=INK)
    if i < 2:
        a = s2.shapes.add_shape(MSO_SHAPE.CHEVRON, Pt(x + gw + 4), Pt(TOP + 32), Pt(10), Pt(16))
        a.fill.solid(); a.fill.fore_color.rgb = MUTED; a.line.fill.background()

# B. Solution overview (template pointers) + evidence
by = TOP + 88
ptr(s2, L, by, 560, "Proposed Solution (Describe your Idea/Solution/Prototype)", size=10.5)
blocks = [
    ("Detailed explanation of the proposed solution", [
        [("Input: ", {"bold": True}), ("Sentinel-2 L2A, 10 bands (B2–B8A, B11, B12), chosen by cloud-aware search or uploaded as a GeoTIFF.", {})],
        [("Model: ", {"bold": True}), ("deep SR backbone + ", {}), ("low-frequency hard constraint", {"bold": True}), (" (FFT), so the network can only add detail, never change measured radiometry.", {})],
        [("Output: ", {"bold": True}), ("2.5 m 10-band GeoTIFF + σ-uncertainty GeoTIFF + validation metrics, in the source UTM grid.", {})],
    ]),
    ("How it addresses the problem", [
        [("Fine-scale interpretation: ", {"bold": True}), ("roofs, stadiums, bunds and debris tracks resolved (1.17× edge energy vs bicubic).", {})],
        [("Consistency by construction: ", {"bold": True}), ("degrading the output reproduces the input at ", {}), ("51.1 dB (10 m) / 36.8 dB (20 m)", {"bold": True}), (", spectral angle 0.48 degrees.", {})],
        [("Uncertainty handled: ", {"bold": True}), ("inferred detail is flagged pixel by pixel; mean σ = 0.0008 reflectance.", {})],
    ]),
]
yy = by + 20
for hd, items in blocks:
    text(s2, L, yy, 565, 14, [hd], size=10, bold=True, color=BLUE)
    text(s2, L + 2, yy + 16, 563, 84, [{"runs": it, "bullet": True, "indent": 9, "space_after": 3.5} for it in items], size=9, color=INK, line=1.05)
    yy += 103
# evidence panel
ex, ew = 606, R - 606
ev = box(s2, ex, by, ew, 206, fill=TINT, radius=0.04)
text(s2, ex + 10, by + 6, ew - 20, 12, ["Prototype output · M. Chinnaswamy Stadium, Bengaluru · 13 Mar 2024"], size=7.3, bold=True, color=NAVY)
iw = (ew - 20 - 12) / 3
for i, (f, lab) in enumerate([("blr_10m.jpg", "Sentinel-2 · 10 m"), ("blr_bic.jpg", "Bicubic · 2.5 m"), ("blr_sr.jpg", "TRINETRA-SR · 2.5 m")]):
    px = ex + 10 + i * (iw + 6)
    picture(s2, A / f, px, by + 22, iw, iw)
    text(s2, px, by + 23 + iw, iw, 11, [lab], size=6.8, bold=(i == 2), color=GREEN if i == 2 else MUTED, align=PP_ALIGN.CENTER)
stats = [("4×", "resolution"), ("10", "bands"), ("51 dB", "consistency"), ("19 MB", "model")]
sw = (ew - 20 - 3 * 5) / 4
for i, (big, small) in enumerate(stats):
    cx = ex + 10 + i * (sw + 5); cy = by + 40 + iw
    box(s2, cx, cy, sw, 40, fill=WHITE, radius=0.1)
    text(s2, cx, cy + 3, sw, 20, [big], size=13, bold=True, color=SAFF if i == 2 else NAVY, align=PP_ALIGN.CENTER)
    text(s2, cx, cy + 23, sw, 12, [small], size=6.8, color=MUTED, align=PP_ALIGN.CENTER)
ev.height = Pt(cy + 40 + 10 - by)

# C. Innovation chips
cy = by + 212
ptr(s2, L, cy, 600, "Innovation and uniqueness of the solution", size=10.5)
chips = [("shield", BLUE, "Trust map", "σ layer + validation on every product"),
         ("wave", NAVY, "Physics-consistent", "FFT hard constraint keeps radiometry"),
         ("layers", GREEN, "All-band 2.5 m", "NDVI · NDRE · NBR at 2.5 m"),
         ("laptop", SAFF, "Zero-server", "19 MB, runs offline / air-gapped"),
         ("clock", MUTED, "Multi-temporal + conformal", "observed detail, guaranteed intervals (v2)")]
chw = (R - L - 4 * 6) / 5
for i, (ic, col, hd, sub) in enumerate(chips):
    x = L + i * (chw + 6)
    box(s2, x, cy + 18, chw, 50, fill=TINT, radius=0.12)
    icon(s2, ic, x + 6, cy + 31, 24, fill=col)
    text(s2, x + 35, cy + 20, chw - 38, 46, [[(hd, {"bold": True, "color": NAVY, "size": 8.8})], [(sub, {"color": MUTED, "size": 7.3})]], anchor=MSO_ANCHOR.MIDDLE)

# D. Tech stack strip
ty = BOTTOM - 14
text(s2, L, ty, 60, 14, ["Tech stack"], size=8, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
xx = L + 58
for t_ in ["Python", "PyTorch", "ONNX Runtime Web", "WASM / WebGPU", "Rasterio · GDAL", "STAC · COG", "Three.js", "Leaflet", "FastAPI · Docker", "GitHub Pages"]:
    w = pill(s2, xx, ty + 1, t_, BLUE_T, color=NAVY, size=6.8, h=12)
    xx += w + 4
notes(s2, "Problem to gap to solution. Numbers are measured on five Indian study areas; the images are real prototype output.")

# ================================================================== 3 · TECHNICAL APPROACH
team_oval(s3); remove(body_box(s3))
ptr(s3, L, TOP - 4, 900, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", size=10.5)

# End-to-end flowchart (snake)
fx0, fy = L, TOP + 16
box(s3, fx0, fy, 528, 158, fill=TINT, radius=0.03)
head(s3, fx0 + 8, fy + 6, 300, "End-to-end flowchart", "gears", BLUE, size=9)
nw, nh, gap = 94, 40, 10
r1 = [("① Input", "point + dates\nor GeoTIFF"), ("② STAC search", "Earth Search\nL2A, <20% cloud"), ("③ Cloud scoring", "SCL mask\nclear ≥ 97%"),
      ("④ COG read", "window only\n10 bands"), ("⑤ Harmonise", "baseline-04\n÷10 000, stack")]
r2 = [("⑩ Deliver", "GeoTIFF · σ\nviewer"), ("⑨ Validate", "PSNR · SAM\nΔNDVI"), ("⑧ 8× TTA", "mean + σ\nper pixel"),
      ("⑦ SR ×4", "ONNX + FFT\nconstraint"), ("⑥ Tile", "128 px, 32 px\nfeathered overlap")]
xs = [fx0 + 10 + i * (nw + gap) for i in range(5)]
y1, y2 = fy + 30, fy + 104
for i, (t_, sub) in enumerate(r1):
    node(s3, xs[i], y1, nw, nh, t_, sub, fill=WHITE, line=LINE if i != 2 else SAFF)
    if i < 4:
        arrow(s3, xs[i] + nw, y1 + nh / 2, xs[i + 1], y1 + nh / 2, color=BLUE, width=1.2)
for i, (t_, sub) in enumerate(r2):
    hl = i in (1, 3)
    node(s3, xs[i], y2, nw, nh, t_, sub, fill=(GREEN_T if i == 0 else WHITE), line=(GREEN if hl else LINE))
    if i < 4:
        arrow(s3, xs[i + 1], y2 + nh / 2, xs[i] + nw, y2 + nh / 2, color=BLUE, width=1.2)
arrow(s3, xs[4] + nw / 2, y1 + nh, xs[4] + nw / 2, y2, color=BLUE, width=1.2)
text(s3, xs[2] - 6, y1 + nh + 4, nw + 12, 26, [[("retry next scene if cloudy", {"italic": True, "color": SAFF, "size": 6.5})]], align=PP_ALIGN.CENTER)

# System architecture
ax0 = fx0 + 536
aw = R - ax0
box(s3, ax0, fy, aw, 158, fill=TINT, radius=0.03)
head(s3, ax0 + 8, fy + 6, 300, "System architecture", "server", NAVY, size=9)
bx, bw_ = ax0 + 8, 220
b = box(s3, bx, fy + 28, bw_, 124, fill=WHITE, line=NAVY, radius=0.04)
text(s3, bx + 6, fy + 30, bw_ - 12, 11, [[("BROWSER · client-side engine", {"bold": True, "size": 6.8, "color": NAVY})]])
layers = [("Presentation", "Three.js globe · Leaflet map · swipe viewer", BLUE_T),
          ("Data engine", "STAC client · SCL scoring · COG reader · proj4", GREEN_T),
          ("Inference", "ONNX Runtime Web · WASM / WebGPU · 19 MB model", SAFF_T),
          ("Products", "TTA σ · metrics · GeoTIFF writer · renderer", TINT)]
for i, (ln, desc, col) in enumerate(layers):
    ly = fy + 44 + i * 26.5
    box(s3, bx + 6, ly, bw_ - 12, 23, fill=col, radius=0.12)
    text(s3, bx + 12, ly + 1, bw_ - 24, 21, [[(ln, {"bold": True, "size": 7.2, "color": NAVY})], [(desc, {"size": 6.3, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE)
ext = ["Earth Search STAC", "AWS S3 · S2 COGs", "Photon / OSM search", "GitHub Pages (static)", "Python batch → ONNX"]
ex0 = bx + bw_ + 14
for i, e in enumerate(ext):
    ey = fy + 30 + i * 24.5
    node(s3, ex0, ey, R - ex0 - 8, 19, e, None, fill=WHITE, line=LINE, tsize=6.8, dash=(i == 4))
    arrow(s3, ex0 - 1, ey + 9.5, bx + bw_ + 1, ey + 9.5, color=MUTED, width=0.9)

# Bottom: tech stack | I/O | deployment | features | links
py = fy + 166
ptr(s3, L, py, 900, "Technologies to be used (e.g. programming languages, frameworks, hardware)", size=10.5)
cy0 = py + 18
cw5 = (R - L - 4 * 7) / 5
ch = BOTTOM - cy0
cols = [(L + i * (cw5 + 7)) for i in range(5)]
for x in cols:
    box(s3, x, cy0, cw5, ch, fill=TINT, radius=0.05)
# tech stack
head(s3, cols[0] + 6, cy0 + 6, cw5, "Tech stack", "code", BLUE, size=8.5)
stack = [("AI", "Python · PyTorch · SEN2SR-Lite · ONNX"), ("Geo", "Rasterio · GDAL · STAC · COG · proj4"),
         ("Web", "JavaScript · ONNX Runtime Web · Three.js · Leaflet · geotiff.js"), ("Ops", "GitHub Pages · FastAPI · Docker"),
         ("HW", "Any laptop CPU; GPU optional (training / batch)")]
text(s3, cols[0] + 6, cy0 + 26, cw5 - 12, ch - 30, [{"runs": [(k + "  ", {"bold": True, "color": BLUE}), (v, {})], "space_after": 3} for k, v in stack], size=7, color=INK)
# I/O
head(s3, cols[1] + 6, cy0 + 6, cw5, "Input / Output", "upload", GREEN, size=8.5)
node(s3, cols[1] + 6, cy0 + 26, cw5 - 12, 64, "INPUT",
     "S2 L2A · 10 bands\n4 × 10 m + 6 × 20 m\nAOI point + dates, or\nGeoTIFF (10/12/13 bands)", fill=WHITE, tcol=BLUE, tsize=7, ssize=6.6)
arrow(s3, cols[1] + cw5 / 2, cy0 + 91, cols[1] + cw5 / 2, cy0 + 99, color=MUTED)
node(s3, cols[1] + 6, cy0 + 100, cw5 - 12, 64, "OUTPUT",
     "2.5 m GeoTIFF · 10 bands\nσ-uncertainty GeoTIFF\nmetrics JSON · viewer\nsource UTM, reflectance ×10⁴", fill=WHITE, tcol=GREEN, tsize=7, ssize=6.6)
# deployment
head(s3, cols[2] + 6, cy0 + 6, cw5, "Deployment", "rocket", SAFF, size=8.5)
deps = [("Web app (client-side)", "GitHub Pages; nothing uploaded", BUILT), ("Offline / air-gapped", "serve docs/ + internal COG mirror", BUILT),
        ("On-prem API", "FastAPI + Docker", BUILT), ("Batch production", "Python tiler, GPU optional", BUILT), ("NTRO integration", "STAC mirror, QGIS plug-in", PLAN)]
for i, (a_, b_, st) in enumerate(deps):
    yy = cy0 + 26 + i * 27
    text(s3, cols[2] + 6, yy, cw5 - 46, 26, [[(a_, {"bold": True, "size": 7, "color": NAVY})], [(b_, {"size": 6.3, "color": MUTED})]])
    status(s3, cols[2] + cw5 - 5, yy + 2, st)
# features
head(s3, cols[3] + 6, cy0 + 6, cw5, "Features", "sliders", NAVY, size=8.5)
feats = ["Search any place · use my location", "Cloud-aware scene selection", "Swipe viewer · zoom · full screen", "NDVI 10 m vs 2.5 m",
         "Uncertainty overlay", "Auto validation metrics", "GeoTIFF export (3 layers)", "Upload own GeoTIFF", "3-D globe · 5 study areas"]
text(s3, cols[3] + 6, cy0 + 26, cw5 - 10, ch - 30, [{"runs": f, "bullet": True, "char": "✓", "bullet_color": GREEN, "indent": 9, "space_after": 1.6} for f in feats], size=7, color=INK)
# links
head(s3, cols[4] + 6, cy0 + 6, cw5, "Links", "link", BLUE, size=8.5)
s3.shapes.add_picture(str(A / "qr_demo.png"), Pt(cols[4] + 8), Pt(cy0 + 27), Pt(52), Pt(52))
text(s3, cols[4] + 64, cy0 + 30, cw5 - 68, 48, [[("Web app", {"bold": True, "color": NAVY, "size": 7.5})], [("scan or open the link below", {"color": MUTED, "size": 6.3})]])
links = [("globe", "Web app", DEMO), ("github", "GitHub", CODE), ("book", "Technical design", "TECHNICAL_DESIGN.md"),
         ("youtube", "Demo video", YOUTUBE or "[add YouTube link]")]
for i, (ic, lab, url) in enumerate(links):
    yy = cy0 + 84 + i * 21
    icon(s3, ic, cols[4] + 6, yy + 1, 14, fill=RED if ic == "youtube" else NAVY)
    text(s3, cols[4] + 24, yy, cw5 - 28, 20, [[(lab, {"bold": True, "size": 6.8, "color": NAVY})], [(url, {"size": 5.6, "color": SAFF if url.startswith("[") else BLUE})]])
notes(s3, "Flowchart and architecture reflect the running prototype. v2 items are marked.")

# ================================================================== 4 · INNOVATION + IMPLEMENTATION
team_oval(s4); remove(body_box(s4))
title(s4, "INNOVATION & IMPLEMENTATION", 28)

# model architecture (research style)
my = TOP
box(s4, L, my, R - L, 122, fill=TINT, radius=0.03)
head(s4, L + 8, my + 5, 300, "Model architecture · TRINETRA-SR", "brain", NAVY, size=9)
lx = R - 330
for i, (lab, col) in enumerate([("Fidelity", BLUE), ("Detail", SAFF), ("Trust", GREEN)]):
    box(s4, lx + i * 70, my + 9, 9, 9, fill=col, radius=0.3)
    text(s4, lx + i * 70 + 12, my + 6, 56, 14, [lab], size=7, color=INK, anchor=MSO_ANCHOR.MIDDLE)
b = box(s4, lx + 212, my + 9, 14, 9, fill=WHITE, radius=0.2); dashed(b, MUTED)
text(s4, lx + 230, my + 6, 100, 14, ["dashed = v2 (planned)"], size=7, color=INK, anchor=MSO_ANCHOR.MIDDLE)
ry = my + 30
tensor(s4, L + 8, ry + 4, 88, 50, WHITE, NAVY, "x : K×10×H×W", "Sentinel-2 L2A\nK dates · 10 m")
chain = [("Shared encoder", "Conv 3×3 ×2\nper date", BLUE, BLUE_T, False, 82),
         ("Temporal attention", "cross-date fusion\nsub-pixel shifts", SAFF, SAFF_T, True, 94),
         ("SR backbone", "N × SPAN / Swin blocks\n1.5 M → ≤15 M params", SAFF, SAFF_T, False, 118),
         ("Pixel-shuffle ×4", "sub-pixel conv\nupsampler", SAFF, SAFF_T, False, 84),
         ("LF hard constraint", "ŷ = y + F⁻¹[M ⊙ F(U(x) − y)]\nPSF/MTF-matched M (v2)", BLUE, BLUE_T, False, 152)]
cx = L + 108
nodes_x = []
for i, (t_, sub, col, tint, dsh, w) in enumerate(chain):
    node(s4, cx, ry + 6, w, 44, t_, sub, fill=tint, line=col, dash=dsh, tsize=7.3, ssize=6.2)
    nodes_x.append((cx, w))
    arrow(s4, cx - 10, ry + 28, cx, ry + 28, color=MUTED, width=1)
    cx += w + 10
tensor(s4, cx, ry + 4, 96, 50, GREEN_T, GREEN, "ŷ : 10×4H×4W", "2.5 m · 10 bands")
arrow(s4, cx - 10, ry + 28, cx, ry + 28, color=MUTED, width=1)
# skip connection U(x)
k = s4.shapes.add_connector(1, Pt(L + 50), Pt(ry), Pt(nodes_x[4][0] + 30), Pt(ry))
k.line.color.rgb = BLUE; k.line.width = Pt(0.9); k.line.dash_style = MSO_LINE.DASH
text(s4, L + 170, ry - 11, 300, 10, [[("skip: U(x), antialiased bicubic ×4 (low frequencies of the measurement)", {"italic": True, "size": 6.3, "color": BLUE})]])
# branch row
by2 = ry + 60
bx2, tx_ = nodes_x[2][0], nodes_x[4][0] + 30
node(s4, bx2, by2, 124, 26, "Diffusion / flow refiner", "4–15 steps · high-freq residual r", fill=SAFF_T, line=SAFF, dash=True, tsize=7, ssize=6)
node(s4, bx2 + 132, by2, 112, 26, "Fidelity ↔ Detail dial", "ŷλ = ŷ + λ·r ,  λ ∈ [0,1]", fill=SAFF_T, line=SAFF, dash=True, tsize=7, ssize=6)
node(s4, tx_, by2, 112, 26, "8× dihedral TTA", "σ = std of inverted outputs", fill=GREEN_T, line=GREEN, tsize=7, ssize=6)
node(s4, tx_ + 122, by2, 126, 26, "Conformal calibration", "P(y ∈ ŷ ± q̂σ) ≥ 1 − α", fill=GREEN_T, line=GREEN, dash=True, tsize=7, ssize=6)
arrow(s4, tx_ + 112, by2 + 13, tx_ + 122, by2 + 13, color=GREEN, width=1)
arrow(s4, bx2 + 124, by2 + 13, bx2 + 132, by2 + 13, color=SAFF, width=1)
arrow(s4, bx2 + 59, ry + 50, bx2 + 59, by2, color=SAFF, width=1)
arrow(s4, cx + 44, ry + 54, tx_ + 56, by2, color=GREEN, width=1)

# three pillars
py4 = my + 130
pw = (R - L - 2 * 8) / 3
pillars = [
    ("bullseye", BLUE, BLUE_T, "1 · FIDELITY", "measured information is never altered", [
        ("4× super-resolution", "10 m → 2.5 m ground sampling, 16× pixels", BUILT),
        ("Low-frequency constraint", "FFT hard layer restores the input's low band", BUILT),
        ("PSF / MTF consistency", "per-band Gaussian PSF from the S2 MTF spec", PLAN),
        ("Spectral consistency", "SAM 0.48° (10 m) · ΔNDVI 0.007 measured", BUILT),
        ("All-band reconstruction", "B2–B8A, B11, B12 to 2.5 m, SWIR via fusion", BUILT)]),
    ("zoom", SAFF, SAFF_T, "2 · DETAIL", "recover structure a 10 m pixel hides", [
        ("Deep SR backbone", "CNN + SPAN attention (v1); Swin / Mamba ≤15 M (v2)", BUILT),
        ("Temporal attention", "fuse 4–8 revisits; sub-pixel shifts add real detail", PLAN),
        ("High-frequency reconstruction", "residual learning over U(x); 1.17× edge energy", BUILT),
        ("Diffusion / flow refinement", "optional residual refiner, 4–15 steps", PLAN),
        ("Fidelity ↔ Detail control", "user dial λ blends fidelity and refiner outputs", PLAN)]),
    ("shield", GREEN, GREEN_T, "3 · TRUST", "every pixel says how sure it is", [
        ("Uncertainty estimation", "per-band σ layer shipped as GeoTIFF", BUILT),
        ("TTA / ensemble", "8 rotations and flips; σ = spread of outputs", BUILT),
        ("Conformal calibration", "split-conformal q̂ → intervals with ≥ 1−α coverage", PLAN),
        ("Hallucination detection", "phantom-object audit vs HR; must fall in high σ", PLAN),
        ("Observed vs inferred detail", "multi-date agreement map per pixel", PLAN)]),
]
for i, (ic, col, tint, hd, sub, items) in enumerate(pillars):
    x = L + i * (pw + 8)
    box(s4, x, py4, pw, 190, fill=WHITE, line=LINE, radius=0.04)
    box(s4, x, py4, pw, 26, fill=tint, radius=0.15)
    icon(s4, ic, x + 6, py4 + 4, 18, fill=col)
    text(s4, x + 29, py4 + 2, pw - 34, 22, [[(hd, {"bold": True, "color": col, "size": 9})], [(sub, {"color": MUTED, "size": 6.5})]], anchor=MSO_ANCHOR.MIDDLE)
    for j, (nm, ds, st) in enumerate(items):
        yy = py4 + 31 + j * 31.5
        text(s4, x + 8, yy, pw - 52, 30, [[(nm, {"bold": True, "size": 7.8, "color": NAVY})], [(ds, {"size": 6.6, "color": INK})]])
        status(s4, x + pw - 6, yy + 2, st)
        if j < 4:
            ln = s4.shapes.add_connector(1, Pt(x + 8), Pt(yy + 29.5), Pt(x + pw - 8), Pt(yy + 29.5))
            ln.line.color.rgb = LINE; ln.line.width = Pt(0.5)

# innovation + roadmap
iy = py4 + 197
iw4 = 372
box(s4, L, iy, iw4, BOTTOM - iy, fill=DARK, radius=0.05)
text(s4, L + 10, iy + 5, iw4 - 20, 14, [[("4 · INNOVATION", {"bold": True, "color": RGBColor(0xFF, 0xB4, 0x54)})]], size=9)
inn = [("First trust-first SR pipeline for Sentinel-2 in India: ", "fidelity, detail and uncertainty as separate, auditable layers."),
       ("Exact ONNX reformulation of the FFT constraint: ", "a 19 MB model with 9×10⁻⁷ parity, runs in any browser."),
       ("Observed vs inferred: ", "multi-date fusion + conformal intervals turn SR into evidence, not decoration."),
       ("Sovereign & offline: ", "no server, no foreign VHR dependency; air-gapped ready.")]
text(s4, L + 10, iy + 20, iw4 - 20, BOTTOM - iy - 22, [{"runs": [(a_, {"bold": True, "color": WHITE}), (b_, {"color": RGBColor(0xC9, 0xD6, 0xEA)})], "bullet": True, "bullet_color": RGBColor(0xFF, 0xB4, 0x54), "indent": 8, "space_after": 1.5} for a_, b_ in inn], size=6.9)
rx4 = L + iw4 + 8
rw4 = R - rx4
box(s4, rx4, iy, rw4, BOTTOM - iy, fill=TINT, radius=0.05)
text(s4, rx4 + 10, iy + 5, rw4, 14, [[("5 · IMPLEMENTATION ROADMAP", {"bold": True, "color": NAVY})]], size=9)
phases = [("Now", "v1 prototype", "live web app, 5 sites", GREEN, True), ("Wk 1–3", "Data", "SEN2NAIPv2 · WorldStrat\nIndia ref. set", BLUE, False),
          ("Wk 3–6", "Fine-tune", "PSF/MTF loss\nIndian adaptation", BLUE, False), ("Wk 6–9", "Temporal + refiner", "attention fusion\nλ dial", SAFF, False),
          ("Wk 8–10", "Trust v2", "conformal · phantom\naudit · OpenSR-test", GREEN, False), ("Wk 10–12", "Finale", "task validation\nQGIS · hardening", NAVY, False)]
tx0, tx1 = rx4 + 30, R - 30
tly = iy + 44
ln = s4.shapes.add_connector(1, Pt(tx0), Pt(tly), Pt(tx1), Pt(tly)); ln.line.color.rgb = MUTED; ln.line.width = Pt(1.5)
step = (tx1 - tx0) / (len(phases) - 1)
for i, (wk, nm, ds, col, done) in enumerate(phases):
    px_ = tx0 + i * step
    d = box(s4, px_ - 6, tly - 6, 12, 12, fill=col if done else WHITE, line=col, shape=MSO_SHAPE.OVAL)
    d.line.width = Pt(2)
    text(s4, px_ - 40, tly - 22, 80, 12, [[(wk, {"bold": True, "size": 7, "color": col})]], align=PP_ALIGN.CENTER)
    text(s4, px_ - 42, tly + 8, 84, 40, [[(nm, {"bold": True, "size": 7, "color": NAVY})]] + [[(l_, {"size": 6, "color": MUTED})] for l_ in ds.split("\n")], align=PP_ALIGN.CENTER)
notes(s4, "Diagram: solid = running in v1, dashed = planned v2. Pillars map 1:1 to the PS: fidelity, detail, uncertainty.")

# ================================================================== 5 · FEASIBILITY, IMPACT & VALIDATION
team_oval(s5); remove(body_box(s5))
title(s5, "FEASIBILITY, IMPACT & VALIDATION", 26)
hw = (R - L - 10) / 2
# risks table
ptr(s5, L, TOP - 2, hw, "Potential challenges and risks  ·  Strategies for overcoming these challenges", size=9)
risks = [
    ("Hallucinated detail", [("HIGH", {"bold": True, "color": RED})], "FFT hard constraint + σ map; flagged, never hidden; phantom audit (v2)"),
    ("Domain gap (US-trained)", [("HIGH", {"bold": True, "color": RED})], "Fine-tune on Indian refs (Cartosat via NRSC/NTRO, WorldStrat)"),
    ("Clouds / monsoon", [("MED", {"bold": True, "color": SAFF})], "12-month SCL-scored search; multi-date fusion; SAR guidance (v2)"),
    ("No 2.5 m ground truth", [("MED", {"bold": True, "color": SAFF})], "Self-consistency now; OpenSR-test + task-level tests (v2)"),
    ("Cross-sensor misalignment", [("MED", {"bold": True, "color": SAFF})], "Co-registration + shift-tolerant loss in training"),
    ("Scale & compute", [("LOW", {"bold": True, "color": GREEN})], "Browser for on-demand; tiled GPU batch for state-scale mosaics"),
]
table(s5, L, TOP + 16, hw, [104, 34, hw - 138], risks, ["Risk", "Level", "Mitigation"], size=6.8, row_h=21)
# feasibility 2x2
fx = L + hw + 10
ptr(s5, fx, TOP - 2, hw, "Analysis of the feasibility of the idea", size=9)
fe = [("gears", BLUE, "Technical", ["Prototype live; 5 Indian sites validated", "19 MB model · ≈6 s per 1.28 km tile, laptop CPU"]),
      ("rupee", SAFF, "Economic", ["Data free (Copernicus), 5-day revisit", "Hosting ₹0 (static); training on free GPUs"]),
      ("people", GREEN, "Operational", ["No install: open a link, or run air-gapped", "Standard GeoTIFF output for GIS workflows"]),
      ("scale", NAVY, "Legal & data", ["Model CC0-1.0; Sentinel data open licence", "No personal data; nothing leaves the device"])]
fw_ = (hw - 8) / 2
for i, (ic, col, hd, lines) in enumerate(fe):
    x = fx + (i % 2) * (fw_ + 8); y = TOP + 16 + (i // 2) * 72
    box(s5, x, y, fw_, 66, fill=TINT, radius=0.08)
    icon(s5, ic, x + 7, y + 7, 20, fill=col)
    text(s5, x + 32, y + 7, fw_ - 80, 20, [hd], size=9, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    pill(s5, x + fw_ - 44, y + 11, "HIGH", GREEN, w=38)
    text(s5, x + 8, y + 31, fw_ - 14, 34, [{"runs": l_, "bullet": True, "indent": 8, "space_after": 1} for l_ in lines], size=6.9, color=INK)
# validation table
vy = TOP + 168
head(s5, L, vy, hw, "Validation metrics · v1 measured on 5 Indian study areas (TTA = 8)", "chart", NAVY, size=8.5)
vm = [
    ("Consistency PSNR, 10 m bands", [("51.1 dB", {"bold": True})], "49.2 – 53.1", "> 40 dB"),
    ("Consistency PSNR, 20 m bands", [("36.8 dB", {"bold": True})], "34.7 – 40.0", "> 33 dB"),
    ("Spectral angle (SAM), 10 m", [("0.48°", {"bold": True})], "0.23 – 0.62", "< 1.5°"),
    ("NDVI deviation (MAE)", [("0.0066", {"bold": True})], "0.003 – 0.009", "< 0.02"),
    ("Edge-energy gain vs bicubic", [("1.17×", {"bold": True})], "1.14 – 1.21", "> 1.1×"),
    ("Pixels flagged σ > 0.01", [("0.05 %", {"bold": True})], "0 – 0.19", "< 5 %"),
    ("Browser vs Python parity", [("0.02 dB", {"bold": True})], "same scene", "< 0.1 dB"),
]
table(s5, L, vy + 20, hw, [150, 62, 72, hw - 284], vm, ["Metric", "Mean", "Range", "Target"], size=6.8, row_h=18.2)
text(s5, L, vy + 20 + 8 * 18.2 + 3, hw, 22, [[("v2 reference validation: ", {"bold": True, "color": NAVY}),
     ("OpenSR-test and SEN2NAIPv2 cross-sensor pairs (PSNR · SSIM · LPIPS · SAM), India Cartosat set, and building / road / field-boundary / change-detection F1.", {"color": MUTED})]], size=6.4)
# impact
ix = fx
ptr(s5, ix, vy - 2, hw, "Potential impact on the target audience", size=9)
apps = [("lock2", BLUE, "Defence & intelligence", "site-level look without VHR tasking"), ("flood", SAFF, "Disaster management", "landslide / flood extent at 2.5 m"),
        ("wheat", GREEN, "Agriculture", "field-level NDVI, smallholder bunds"), ("city", NAVY, "Urban planning", "buildings, roads, encroachment"),
        ("water", BLUE, "Water & environment", "ponds, shorelines, forest loss"), ("grad", MUTED, "Research & education", "open, runs on a student laptop")]
aw_ = (hw - 2 * 6) / 3
for i, (ic, col, hd, ds) in enumerate(apps):
    x = ix + (i % 3) * (aw_ + 6); y = vy + 18 + (i // 3) * 44
    box(s5, x, y, aw_, 39, fill=TINT, radius=0.1)
    icon(s5, ic, x + 5, y + 8, 22, fill=col)
    text(s5, x + 31, y + 2, aw_ - 34, 35, [[(hd, {"bold": True, "size": 7.3, "color": NAVY})], [(ds, {"size": 6.3, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE)
ptr(s5, ix, vy + 108, hw, "Benefits of the solution (social, economic, environmental, etc.)", size=9)
ben = [("people", BLUE, "Social", ["Faster, better-targeted relief", "Field-level farm advisories"]),
       ("rupee", SAFF, "Economic", ["2.5 m-class from free data", "No per-km² VHR licences"]),
       ("leaf", GREEN, "Environmental", ["Finer land & water monitoring", "Early landslide / erosion maps"]),
       ("flag", NAVY, "Strategic", ["Sovereign, offline capability", "No foreign data dependency"])]
bw5 = (hw - 3 * 6) / 4
for i, (ic, col, hd, ds) in enumerate(ben):
    x = ix + i * (bw5 + 6); y = vy + 126
    box(s5, x, y, bw5, BOTTOM - y, fill=WHITE, line=LINE, radius=0.1)
    icon(s5, ic, x + 5, y + 6, 18, fill=col)
    text(s5, x + 27, y + 5, bw5 - 30, 20, [hd], size=7.6, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    text(s5, x + 6, y + 28, bw5 - 10, BOTTOM - y - 30, [{"runs": d_, "bullet": True, "indent": 7, "space_after": 3} for d_ in ds], size=7, color=INK)
notes(s5, "All measured values come from the prototype's automatic validation on five Indian scenes. Targets for v2 are stated separately.")

# ================================================================== 6 · RESEARCH & REFERENCES
team_oval(s6); remove(body_box(s6))
lw6 = 606
ptr(s6, L, TOP - 4, lw6, "Details / Links of the reference and research work", size=10.5)
head(s6, L, TOP + 14, lw6, "Key findings from the literature → our design decisions", "flask", NAVY, size=8.5)
kf = [("SEN2SR [1]: a low-frequency hard constraint gives radiometric consistency; > 15 M parameters brought no significant gain.", "Keep the constraint; compact 1.5–15 M backbone."),
      ("GeoSR-Bench [13]: PSNR / SSIM gains often do not track downstream-task gains, and can even run opposite.", "Validate on buildings, roads, bunds, change."),
      ("LDSR-S2 [6]: the only prior S2 SR with pixel-wise uncertainty, but diffusion is slow.", "TTA σ now; diffusion only as optional refiner."),
      ("SEN2NAIP [9]: 2,851 real + 17,657 synthetic pairs; real pairs are mis-registered.", "Synthetic pre-train, shift-tolerant fine-tune."),
      ("HighRes-net, WorldStrat, MuS2 [8, 10, 11]: multiple revisits add true sub-pixel information.", "Temporal attention front-end in v2.")]
for i, (f_, d_) in enumerate(kf):
    y = TOP + 34 + i * 27
    box(s6, L, y, 392, 24, fill=TINT, radius=0.12)
    text(s6, L + 6, y, 382, 24, [f_], size=6.7, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s6, L + 394, y + 12, L + 408, y + 12, color=SAFF, width=1.2)
    box(s6, L + 410, y, lw6 - 410, 24, fill=GREEN_T, radius=0.12)
    text(s6, L + 416, y, lw6 - 422, 24, [[(d_, {"bold": True})]], size=6.7, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
ry6 = TOP + 34 + 5 * 27 + 4
head(s6, L, ry6, lw6, "References by domain", "book", BLUE, size=8.5)
dom = [
    ("Super-resolution methods", [
        "[1] Aybar et al., 2025. SEN2SR, radiometrically & spatially consistent SR for Sentinel-2. Remote Sens. Environ.",
        "[2] Wan et al., 2024. SPAN: Swift Parameter-free Attention Network. CVPRW.",
        "[3] Liang et al., 2021. SwinIR. ICCVW.", "[4] Wang et al., 2018. ESRGAN. ECCVW.",
        "[5] Yue et al., 2023. ResShift, efficient diffusion SR. NeurIPS.",
        "[6] Donike et al., 2025. Trustworthy S2 SR with latent diffusion.", "[7] DiffFuSR, 2025. All-band S2 SR. arXiv 2506.11764."]),
    ("Multi-image, data & benchmarks", [
        "[8] Deudon et al., 2020. HighRes-net. arXiv 2002.06460.", "[9] Aybar et al., 2024. SEN2NAIP. Scientific Data.",
        "[10] Cornebise et al., 2022. WorldStrat. NeurIPS D&B.", "[11] MuS2, 2023. Multi-image S2 SR benchmark. Scientific Data.",
        "[12] ESA OpenSR. opensr-test benchmark (GitHub).", "[13] GeoSR-Bench, 2026. Downstream-task SR benchmark. arXiv 2605.00310."]),
    ("Uncertainty & data / tools", [
        "[14] Lakshminarayanan et al., 2017. Deep ensembles. NeurIPS.",
        "[15] Wang et al., 2019. Test-time augmentation uncertainty. Neurocomputing.",
        "[16] Angelopoulos & Bates, 2021. Conformal prediction. arXiv 2107.07511.",
        "[17] ESA. Sentinel-2 MSI user guide (bands, MTF).", "[18] Copernicus Data Space · Element 84 Earth Search.",
        "[19] ONNX Runtime Web (WASM / WebGPU)."]),
]
dw = (lw6 - 2 * 6) / 3
for i, (hd, items) in enumerate(dom):
    x = L + i * (dw + 6); y = ry6 + 20
    box(s6, x, y, dw, BOTTOM - y, fill=WHITE, line=LINE, radius=0.04)
    text(s6, x + 6, y + 4, dw - 12, 12, [hd], size=7.4, bold=True, color=BLUE)
    text(s6, x + 6, y + 18, dw - 10, BOTTOM - y - 20, [{"runs": it, "space_after": 3.2} for it in items], size=6.9, color=INK)
# resources panel
rx6 = L + lw6 + 10
rw6 = R - rx6
box(s6, rx6, TOP + 14, rw6, BOTTOM - TOP - 14, fill=DARK, radius=0.04)
text(s6, rx6 + 12, TOP + 20, rw6 - 24, 14, [[("PROJECT RESOURCES", {"bold": True, "color": RGBColor(0xFF, 0xB4, 0x54)})]], size=9)
res = [("qr_demo.png", "globe", "Live web app", DEMO), ("qr_code.png", "github", "Source code (GitHub)", CODE)]
for i, (qr, ic, hd, url) in enumerate(res):
    y = TOP + 40 + i * 92
    box(s6, rx6 + 12, y, 82, 82, fill=WHITE, radius=0.06)
    s6.shapes.add_picture(str(A / qr), Pt(rx6 + 16), Pt(y + 4), Pt(74), Pt(74))
    text(s6, rx6 + 102, y + 14, rw6 - 112, 60, [[(hd, {"bold": True, "color": WHITE, "size": 9})], [(url, {"color": RGBColor(0x7E, 0xC8, 0xF0), "size": 6.4})]])
more = [("book", "Technical design document", "TECHNICAL_DESIGN.md in the repository"), ("youtube", "Demo video", YOUTUBE or "[add YouTube link]"),
        ("file", "Implementation plan", "IMPLEMENTATION_PLAN.md in the repository")]
for i, (ic, hd, sub) in enumerate(more):
    y = TOP + 230 + i * 34
    icon(s6, ic, rx6 + 12, y + 3, 22, fill=RED if ic == "youtube" else BLUE)
    text(s6, rx6 + 42, y, rw6 - 52, 30, [[(hd, {"bold": True, "color": WHITE, "size": 8})], [(sub, {"color": RGBColor(0xFF, 0xB4, 0x54) if sub.startswith("[") else RGBColor(0xC9, 0xD6, 0xEA), "size": 6.5})]], anchor=MSO_ANCHOR.MIDDLE)
text(s6, rx6 + 12, BOTTOM - 40, rw6 - 24, 34, [[("Model credit: ", {"bold": True, "color": WHITE}), ("SEN2SR-Lite (ESA OpenSR, CC0-1.0) is the v1 baseline. The pipeline, trust layer, validation, ONNX engine and web app are ours.", {"color": RGBColor(0xC9, 0xD6, 0xEA)})]], size=6.3)
notes(s6, "References grouped by domain; each key finding maps to a design decision.")

prs.save(str(OUT))
print("saved", OUT)
