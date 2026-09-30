"""Build the SIH 2026 idea-submission deck from the official template.

python build_deck.py  ->  ppt/Sentinels_W_PS26142_Idea.pptx
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
A, ICON = HERE / "assets", HERE / "assets" / "icons"
TEMPLATE = ROOT / "SIH2026-IDEA-Presentation-Format.pptx"
OUT = HERE / "Sentinels_W_PS26142_Idea.pptx"

TEAM, TEAM_ID = "Sentinels_W", "190286"
DEMO = "uchihadzoro.github.io/Deep-Learning-Super-Resolution-Mapping"
CODE = "github.com/UchihaDZoro/Deep-Learning-Super-Resolution-Mapping"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x00, 0x70, 0xC0)      # template footer blue
SAFF = RGBColor(0xE9, 0x7A, 0x1F)      # SIH saffron
GREEN = RGBColor(0x2E, 0x9E, 0x48)     # SIH green
INK = RGBColor(0x1A, 0x1F, 0x2B)
MUTED = RGBColor(0x55, 0x60, 0x70)
TINT = RGBColor(0xF1, 0xF5, 0xFA)
TINT2 = RGBColor(0xFD, 0xF3, 0xEA)
LINE = RGBColor(0xD5, 0xDE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

L, R = 30, 930          # content margins (pt)
TOP, BOTTOM = 92, 490    # below header art, above the blue footer


# ------------------------------------------------------------------ helpers
def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, line=None, margin=0):
    """paras: list of paragraphs; each paragraph is a str or a list of (text, {overrides}) runs.
    A paragraph may also be a dict {"runs": [...], "bullet": True, "size": .., "space_after": ..}."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(margin)
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        spec = p if isinstance(p, dict) else {"runs": p}
        runs = spec["runs"]
        if isinstance(runs, str):
            runs = [(runs, {})]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = spec.get("align", align)
        para.space_after = Pt(spec.get("space_after", space_after))
        if line:
            para.line_spacing = line
        if spec.get("bullet"):
            _bullet(para, spec.get("indent", 11), spec.get("char", "•"), spec.get("bullet_color", BLUE))
        for t, o in runs:
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", spec.get("size", size)))
            f.bold = o.get("bold", spec.get("bold", bold))
            f.italic = o.get("italic", False)
            f.underline = o.get("underline", False)
            f.color.rgb = o.get("color", spec.get("color", color))
    return tb


def _bullet(para, indent, char, color):
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent))))
    pPr.set("indent", str(-int(Pt(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = buClr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    buClr.append(srgb)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buClr); pPr.append(buFont); pPr.append(buChar)


def box(slide, x, y, w, h, fill=TINT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def icon(slide, name, x, y, d, fill=BLUE):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    pad = d * 0.24
    slide.shapes.add_picture(str(ICON / f"{name}.png"), Pt(x + pad), Pt(y + pad), Pt(d - 2 * pad), Pt(d - 2 * pad))
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    p = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w) if w else None, Pt(h) if h else None)
    if border:
        p.line.color.rgb = LINE; p.line.width = Pt(0.75)
    return p


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def pointer(slide, y, label, size=15):
    """Template section pointer, kept verbatim."""
    return text(slide, L, y, R - L, 22, [[("❖ ", {"color": BLUE}), (label, {"underline": True})]],
                size=size, bold=True, color=NAVY)


def team_oval(slide):
    for sh in slide.shapes:
        if sh.name.startswith("Oval") and sh.has_text_frame and "Team" in sh.text_frame.text:
            tf = sh.text_frame
            p0 = tf.paragraphs[0]
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            for r in list(p0.runs)[1:]:
                r._r.getparent().remove(r._r)
            p0.runs[0].text = TEAM
            p0.runs[0].font.size = Pt(11); p0.runs[0].font.bold = True
            p0.runs[0].font.color.rgb = NAVY
            p0.alignment = PP_ALIGN.CENTER
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Pt(2)


def set_title(slide, t, size=None):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = t
            for extra in r[1:]:
                extra._r.getparent().remove(extra._r)
            if size:
                r[0].font.size = Pt(size)
            return sh


def body_box(slide):
    for sh in slide.shapes:
        if sh.shape_type == 17 and sh.name.startswith("TextBox 8"):
            return sh


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ build
prs = Presentation(str(TEMPLATE))
delete_slide(prs, 6)                       # "Important instructions" slide: template says delete before upload
s1, s2, s3, s4, s5, s6 = prs.slides

# ================================================================== 1. TITLE
tb = next(sh for sh in s1.shapes if sh.name == "TextBox 9")
rows = [
    ("Problem Statement ID – ", "26142"),
    ("Problem Statement Title – ", "Deep Learning Based Super Resolution Mapping (SRM) from Medium Resolution Satellite Imageries"),
    ("Theme – ", "Space Technology"),
    ("PS Category – ", "Software"),
    ("Team ID – ", TEAM_ID),
    ("Team Name – ", TEAM),
]
tf = tb.text_frame
paras = [p for p in tf.paragraphs if p.text.strip()]
for p in list(tf.paragraphs):
    if not p.text.strip():
        p._p.getparent().remove(p._p)
for p, (lab, val) in zip(paras, rows):
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    r0 = runs[0]
    r0.text = lab
    r0.font.size = Pt(17); r0.font.bold = True; r0.font.color.rgb = INK
    r1 = copy.deepcopy(r0._r); r0._r.addnext(r1)
    p.runs[1].text = val
    p.runs[1].font.bold = False
    p.runs[1].font.color.rgb = NAVY if lab.startswith("Problem Statement Title") else INK
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.1
    p.space_before = Pt(0)
    p.space_after = Pt(14)
tb.top, tb.height = Pt(150), Pt(360)
tf.word_wrap = True

# small QR to the live prototype, bottom-right
box(s1, 812, 386, 118, 110, fill=WHITE, line=LINE, radius=0.06)
s1.shapes.add_picture(str(A / "qr_demo.png"), Pt(836), Pt(392), Pt(70), Pt(70))
text(s1, 812, 466, 118, 26, [[("Live prototype", {"bold": True, "color": NAVY})], [("scan to open", {"color": MUTED, "size": 8})]],
     size=9.5, align=PP_ALIGN.CENTER)
notes(s1, "Team Sentinels_W, PS 26142 (NTRO). The QR code opens the working prototype.")

# ================================================================== 2. IDEA
team_oval(s2)
t2 = set_title(s2, "TRINETRA-SR: Trustworthy Sentinel-2 Super-Resolution", size=23)
t2.left, t2.top, t2.width, t2.height = Pt(135), Pt(22), Pt(632), Pt(50)
t2.text_frame.word_wrap = True
t2.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
t2.text_frame.auto_size = None
t2.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
remove(body_box(s2))
pointer(s2, TOP, "Proposed Solution (Describe your Idea/Solution/Prototype)", size=17)

colw = 560
y = TOP + 30
blocks = [
    ("Detailed explanation of the proposed solution", [
        [("A deep-learning framework that raises Sentinel-2 L2A imagery from ", {}), ("10 m to 2.5 m (4×)", {"bold": True}), (" on all 10 bands: visible, NIR, red-edge, SWIR.", {})],
        [("Pipeline: cloud-aware scene selection → pre-processing → SR network with a frequency-domain ", {}), ("hard constraint", {"bold": True}), (" → per-pixel uncertainty → validated GeoTIFF.", {})],
        [("Working prototype: ", {"bold": True}), ("pick any location or upload a GeoTIFF; results in ≈6–20 s, entirely in the browser.", {})],
    ]),
    ("How it addresses the problem", [
        [("Reveals small buildings, narrow roads, field bunds and landslide scars that are lost at 10 m.", {})],
        [("Keeps geospatial and spectral consistency: output degraded to 10 m reproduces the input at ", {}), ("49–54 dB", {"bold": True}), (", spectral angle about 0.5 degrees.", {})],
        [("Separates observed from inferred detail with an uncertainty map, as the PS requires.", {})],
    ]),
    ("Innovation and uniqueness of the solution", [
        [("Trust by design: ", {"bold": True}), ("every product ships with an uncertainty layer and automatic validation metrics.", {})],
        [("All-band 2.5 m: ", {"bold": True}), ("NDVI, NDRE, NBR and water indices at 2.5 m, not just an RGB picture.", {})],
        [("Zero-server, 19 MB model: ", {"bold": True}), ("runs offline on a standard laptop; suitable for air-gapped networks.", {})],
    ]),
]
blk_h = (BOTTOM - y) / 3
for head, items in blocks:
    text(s2, L, y, colw, 16, [head], size=13, bold=True, color=BLUE)
    paras = [{"runs": it, "bullet": True, "space_after": 3} for it in items]
    text(s2, L + 4, y + 20, colw - 4, blk_h - 22, paras, size=11.2, color=INK, line=1.05)
    y += blk_h

# right column: evidence
rx, rw = 612, R - 612
panel = box(s2, rx, TOP + 30, rw, 372, fill=TINT, radius=0.04)
text(s2, rx + 12, TOP + 40, rw - 24, 16, ["Prototype result · M. Chinnaswamy Stadium, Bengaluru"], size=9.5, bold=True, color=NAVY)
iw = (rw - 24 - 12) / 3
for i, (f, lab) in enumerate([("blr_10m.jpg", "Sentinel-2 10 m"), ("blr_bic.jpg", "Bicubic 2.5 m"), ("blr_sr.jpg", "TRINETRA 2.5 m")]):
    px = rx + 12 + i * (iw + 6)
    picture(s2, A / f, px, TOP + 60, iw, iw)
    text(s2, px, TOP + 62 + iw, iw, 14, [lab], size=8.5, bold=(i == 2), color=(GREEN if i == 2 else MUTED), align=PP_ALIGN.CENTER)
stats = [("4×", "finer ground\nsampling"), ("10", "spectral bands\nsuper-resolved"), ("51 dB", "mean radiometric\nconsistency"), ("19 MB", "model, no GPU\nor server needed")]
sy = TOP + 60 + iw + 26
sw = (rw - 24 - 8) / 2
for i, (big, small) in enumerate(stats):
    cx = rx + 12 + (i % 2) * (sw + 8); cy = sy + (i // 2) * 62
    box(s2, cx, cy, sw, 56, fill=WHITE, radius=0.08)
    text(s2, cx + 10, cy + 5, sw - 20, 24, [big], size=19, bold=True, color=SAFF if i == 2 else NAVY)
    text(s2, cx + 10, cy + 29, sw - 20, 24, small.split("\n"), size=8.5, color=MUTED)
qy = sy + 130
s2.shapes.add_picture(str(A / "qr_demo.png"), Pt(rx + 12), Pt(qy), Pt(46), Pt(46))
text(s2, rx + 66, qy + 2, rw - 78, 44, [[("Try it live", {"bold": True, "color": NAVY})], [(DEMO, {"color": BLUE, "size": 7.5})]], size=10)
panel.height = Pt(qy + 58 - (TOP + 30))
notes(s2, "Headline: 4x resolution on all 10 bands, radiometry preserved by construction, uncertainty shipped with every product. The images are real prototype output.")

# ================================================================== 3. TECHNICAL APPROACH
team_oval(s3)
remove(body_box(s3))
pointer(s3, TOP - 4, "Technologies to be used (e.g. programming languages, frameworks, hardware)", size=14)
cards = [
    ("brain", "AI & model", ["Python · PyTorch", "SEN2SR-Lite (CNN + SPAN), ESA OpenSR", "ONNX export (exact FFT reformulation)"]),
    ("sat", "Geospatial data", ["Sentinel-2 L2A (Copernicus)", "STAC: Earth Search, Planetary Computer", "Rasterio / GDAL · GeoTIFF / COG"]),
    ("globe", "Inference & web", ["ONNX Runtime Web (WASM / WebGPU)", "JavaScript · Three.js · Leaflet", "OpenStreetMap place search"]),
    ("laptop", "Hardware & deployment", ["Any laptop CPU, no GPU needed", "Static hosting (GitHub Pages)", "FastAPI + Docker for on-prem / air-gapped"]),
]
cw = (R - L - 3 * 10) / 4
cy = TOP + 22
for i, (ic, head, lines) in enumerate(cards):
    cx = L + i * (cw + 10)
    box(s3, cx, cy, cw, 92, fill=TINT, radius=0.08)
    icon(s3, ic, cx + 10, cy + 10, 26, fill=[BLUE, GREEN, SAFF, NAVY][i])
    text(s3, cx + 42, cy + 14, cw - 50, 18, [head], size=11.5, bold=True, color=NAVY)
    text(s3, cx + 10, cy + 42, cw - 18, 48, [{"runs": l, "bullet": True, "indent": 8, "space_after": 1} for l in lines], size=8.8, color=INK)

pointer(s3, cy + 102, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", size=14)
steps = [
    ("dish", "1 · Acquire", "STAC search for Sentinel-2 L2A; rank scenes by cloud-free fraction (SCL mask)"),
    ("layers", "2 · Pre-process", "Windowed COG read, baseline-04 harmonisation, 10-band stack at 10 m"),
    ("chip", "3 · Super-resolve", "×4 network, tiled with feathered overlap; FFT hard constraint keeps low frequencies"),
    ("chart", "4 · Quantify", "8-way rotate/flip ensemble → per-pixel σ; native-resolution consistency checks"),
    ("file", "5 · Deliver", "2.5 m GeoTIFF in source UTM grid + uncertainty layer + metrics + viewer"),
]
fy = cy + 128
fw = (R - L - 4 * 16) / 5
for i, (ic, head, desc) in enumerate(steps):
    fx = L + i * (fw + 16)
    box(s3, fx, fy, fw, 104, fill=WHITE, line=LINE, radius=0.08)
    icon(s3, ic, fx + (fw - 30) / 2, fy + 8, 30, fill=BLUE)
    text(s3, fx + 6, fy + 42, fw - 12, 16, [head], size=10.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s3, fx + 8, fy + 58, fw - 16, 44, [desc], size=8.3, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s3, fx + fw + 2, fy + 52, fx + fw + 14, fy + 52, color=SAFF, width=1.75)

by = fy + 114
bh = BOTTOM - by
# key techniques
kw = 560
box(s3, L, by, kw, bh, fill=TINT2, radius=0.06)
text(s3, L + 12, by + 10, kw - 24, 16, ["Why the output can be trusted"], size=12, bold=True, color=SAFF)
kt = [
    [("Frequency hard constraint: ", {"bold": True}), ("low frequencies of the 2.5 m output are replaced by the input's, so radiometry cannot drift.", {})],
    [("Ensemble uncertainty: ", {"bold": True}), ("eight rotated/mirrored inferences; their spread flags reconstructed detail pixel by pixel.", {})],
    [("Self-validation: ", {"bold": True}), ("each product is degraded back to 10 m and 20 m and scored (PSNR, spectral angle, ΔNDVI).", {})],
]
text(s3, L + 14, by + 32, kw - 28, bh - 36, [{"runs": k, "bullet": True, "space_after": 6} for k in kt], size=10.3, color=INK)
# prototype
px = L + kw + 12
pw = R - px
ph = pw * 780 / 1420
if ph > bh:
    ph = bh; pw2 = ph * 1420 / 780
else:
    pw2 = pw
picture(s3, A / "hero_crop.jpg", px, by, pw2 - 60, (pw2 - 60) * 780 / 1420)
s3.shapes.add_picture(str(A / "qr_demo.png"), Pt(R - 52), Pt(by), Pt(52), Pt(52))
text(s3, R - 56, by + 54, 60, 30, ["Working", "prototype"], size=7.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
notes(s3, "Five-stage pipeline. The same code exists as a Python batch job and as an in-browser engine; they agree within 0.05 dB.")

# ================================================================== 4. FEASIBILITY
team_oval(s4)
remove(body_box(s4))
pointer(s4, TOP - 4, "Analysis of the feasibility of the idea", size=14)
feas = [("Live", "prototype deployed and\npublicly testable today"), ("≈6 s", "per 1.28 km tile on a\nlaptop CPU (WebAssembly)"),
        ("₹0", "data cost: Sentinel-2 is free,\n5-day revisit, archive since 2017"), ("5 sites", "validated in India:\n49–54 dB, spectral angle 0.2–0.6°")]
fw4 = (R - L - 3 * 10) / 4
for i, (big, small) in enumerate(feas):
    fx = L + i * (fw4 + 10)
    box(s4, fx, TOP + 22, fw4, 66, fill=TINT, radius=0.08)
    text(s4, fx + 12, TOP + 27, fw4 - 24, 26, [big], size=21, bold=True, color=[GREEN, NAVY, SAFF, NAVY][i])
    text(s4, fx + 12, TOP + 55, fw4 - 24, 30, small.split("\n"), size=8.8, color=MUTED)

ty = TOP + 100
pointer(s4, ty, "Potential challenges and risks", size=13)
text(s4, 495, ty, R - 495, 22, [[("❖ ", {"color": GREEN}), ("Strategies for overcoming these challenges", {"underline": True})]], size=13, bold=True, color=NAVY)
risks = [
    ("Hallucinated detail", "Model may invent plausible but wrong structures",
     "Frequency hard constraint + per-pixel uncertainty map: inferred detail is flagged, never hidden"),
    ("Domain gap", "Pre-trained on US aerial pairs; Indian terrain differs",
     "Fine-tune on Indian references (Cartosat via NRSC/NTRO, WorldStrat); benchmark on OpenSR-test"),
    ("Clouds & monsoon", "Few clear scenes June–September",
     "12-month SCL-scored scene search; multi-date fusion and SAR guidance planned for v2"),
    ("No 2.5 m ground truth", "Arbitrary scenes cannot be scored against a reference",
     "Self-consistency on every product now; independent references + task-level tests in v2"),
    ("Scale", "State- or national-scale mosaics",
     "Tiled Python batch pipeline on GPU for bulk runs; browser engine for on-demand analysis"),
]
ry = ty + 26
rh = (BOTTOM - ry - 4 * 6) / 5
for i, (name, why, fix) in enumerate(risks):
    yy = ry + i * (rh + 6)
    box(s4, L, yy, 445, rh, fill=TINT2, radius=0.12)
    icon(s4, "warn", L + 8, yy + (rh - 22) / 2, 22, fill=SAFF)
    text(s4, L + 38, yy + 4, 400, rh - 6, [[(name, {"bold": True, "color": INK})], [(why, {"color": MUTED, "size": 8.8})]], size=10.3, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s4, L + 448, yy + rh / 2, 490, yy + rh / 2, color=MUTED, width=1.25)
    box(s4, 495, yy, R - 495, rh, fill=RGBColor(0xEC, 0xF7, 0xEF), radius=0.12)
    icon(s4, "check", 503, yy + (rh - 22) / 2, 22, fill=GREEN)
    text(s4, 533, yy + 3, R - 533 - 8, rh - 6, [fix], size=9.4, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s4, "Feasibility is demonstrated, not assumed: the prototype is live. Each risk is paired with a concrete mitigation.")

# ================================================================== 5. IMPACT
team_oval(s5)
remove(body_box(s5))
pointer(s5, TOP - 4, "Potential impact on the target audience", size=14)
aud = [
    ("shield", BLUE, "Security & intelligence (NTRO)", "A sharper first look at any site without tasking a very-high-resolution satellite."),
    ("flood", SAFF, "Disaster management (NDMA / SDMAs)", "Landslide scars, flood extents and damaged embankments at 2.5 m within days."),
    ("wheat", GREEN, "Agriculture", "Field-level NDVI and boundaries for smallholder farms, most under one hectare."),
    ("city", NAVY, "Urban & infrastructure", "Building, road and encroachment mapping from free, frequent imagery."),
    ("grad", MUTED, "Research & education", "Open, reproducible pipeline that runs on a student laptop."),
]
ay = TOP + 22
ah = 36
for i, (ic, col, head, desc) in enumerate(aud):
    yy = ay + i * (ah + 4)
    icon(s5, ic, L, yy + 3, 30, fill=col)
    text(s5, L + 40, yy + 1, 520, ah, [[(head, {"bold": True, "color": NAVY})], [(desc, {"color": INK, "size": 9.3})]], size=10.6)
# evidence image pair
ex = 600; ew = R - ex
box(s5, ex, ay - 2, ew, 196, fill=TINT, radius=0.05)
text(s5, ex + 10, ay + 4, ew - 20, 14, ["Wayanad, Kerala · July 2024 debris-flow track"], size=9.2, bold=True, color=NAVY)
iw5 = (ew - 20 - 8) / 2
for i, (f, lab) in enumerate([("way_10m.jpg", "Sentinel-2 · 10 m"), ("way_sr.jpg", "TRINETRA-SR · 2.5 m")]):
    xx = ex + 10 + i * (iw5 + 8)
    picture(s5, A / f, xx, ay + 22, iw5, iw5)
    text(s5, xx, ay + 24 + iw5, iw5, 12, [lab], size=8.3, bold=(i == 1), color=GREEN if i else MUTED, align=PP_ALIGN.CENTER)

by5 = ay + 5 * (ah + 4) + 8
pointer(s5, by5, "Benefits of the solution (social, economic, environmental, etc.)", size=14)
bens = [
    ("people", BLUE, "Social", ["Faster, better-targeted disaster relief and damage assessment", "Field-level crop advisories for smallholder farmers"]),
    ("rupee", SAFF, "Economic", ["2.5 m-class analysis from free Sentinel-2 data", "Avoids per-km² licences for commercial imagery in routine monitoring"]),
    ("leaf", GREEN, "Environmental", ["Finer tracking of water bodies, forest loss and mining", "Early mapping of landslide and erosion scars"]),
    ("flag", NAVY, "Strategic", ["Sovereign capability that runs offline, even air-gapped", "No dependence on foreign commercial data providers"]),
]
bw = (R - L - 3 * 10) / 4
bty = by5 + 26
bh5 = BOTTOM - bty
for i, (ic, col, head, desc) in enumerate(bens):
    xx = L + i * (bw + 10)
    box(s5, xx, bty, bw, bh5, fill=TINT, radius=0.07)
    icon(s5, ic, xx + 10, bty + 10, 26, fill=col)
    text(s5, xx + 42, bty + 14, bw - 50, 18, [head], size=11.5, bold=True, color=NAVY)
    text(s5, xx + 10, bty + 44, bw - 20, bh5 - 48, [{"runs": d, "bullet": True, "indent": 9, "space_after": 4} for d in desc], size=9.8, color=INK)
notes(s5, "Primary users: NTRO analysts and disaster agencies; secondary: agriculture, urban planning, research.")

# ================================================================== 6. REFERENCES
team_oval(s6)
remove(body_box(s6))
pointer(s6, TOP - 4, "Details / Links of the reference and research work", size=14)
refs = [
    ("Aybar et al., 2025", "A radiometrically and spatially consistent super-resolution framework for Sentinel-2 (SEN2SR). Remote Sensing of Environment.", "sciencedirect.com/science/article/pii/S0034425725006261"),
    ("SEN2NAIP, 2024", "A large-scale dataset for Sentinel-2 image super-resolution. Scientific Data.", "nature.com/articles/s41597-024-04214-y"),
    ("OpenSR-test", "Benchmark for real-world Sentinel-2 super-resolution (ESA OpenSR).", "github.com/ESAOpenSR/opensr-test"),
    ("Latent-diffusion SR, 2025", "Trustworthy super-resolution of multispectral Sentinel-2 imagery with latent diffusion (pixel-wise uncertainty).", "semanticscholar.org/paper/27fa48af71d55c671c498649b5a65d57fbed13f4"),
    ("DiffFuSR, 2025", "Super-resolution of all Sentinel-2 multispectral bands using diffusion models.", "arxiv.org/abs/2506.11764"),
    ("MuS2, 2023", "A real-world benchmark for Sentinel-2 multi-image super-resolution. Scientific Data.", "nature.com/articles/s41597-023-02538-9"),
    ("GeoSR-Bench, 2026", "Benchmarking SR models for remote sensing via downstream-task integration.", "arxiv.org/abs/2605.00310"),
    ("Data access", "Copernicus Data Space · Element 84 Earth Search (Sentinel-2 COGs on AWS) · Microsoft Planetary Computer.", "browser.dataspace.copernicus.eu"),
]
rx6 = 640
paras = []
for i, (who, what, link) in enumerate(refs, 1):
    paras.append({"runs": [(f"[{i}] ", {"bold": True, "color": BLUE}), (who + ". ", {"bold": True, "color": NAVY}), (what, {})], "space_after": 0})
    paras.append({"runs": [("       " + link, {"color": BLUE, "size": 8.5})], "space_after": 7})
text(s6, L, TOP + 24, rx6 - L - 16, BOTTOM - TOP - 24, paras, size=10.4, color=INK)
box(s6, rx6, TOP + 22, R - rx6, BOTTOM - TOP - 22, fill=TINT, radius=0.04)
text(s6, rx6 + 14, TOP + 32, R - rx6 - 28, 16, ["Our work"], size=12, bold=True, color=NAVY)
qs = 92
for i, (img, head, url) in enumerate([("qr_demo.png", "Live prototype", DEMO), ("qr_code.png", "Source code", CODE)]):
    yy = TOP + 56 + i * (qs + 34)
    box(s6, rx6 + 14, yy, qs + 8, qs + 8, fill=WHITE, radius=0.06)
    s6.shapes.add_picture(str(A / img), Pt(rx6 + 18), Pt(yy + 4), Pt(qs), Pt(qs))
    text(s6, rx6 + qs + 32, yy + 18, R - rx6 - qs - 44, 60, [[(head, {"bold": True, "color": NAVY, "size": 11})], [(url, {"color": BLUE, "size": 8})]], size=10)
text(s6, rx6 + 14, TOP + 56 + 2 * (qs + 34) - 6, R - rx6 - 28, 60,
     [[("Model: ", {"bold": True}), ("SEN2SR-Lite, ESA OpenSR, CC0-1.0, used as the v1 baseline network.", {})],
      [("Ours: ", {"bold": True}), ("pipeline, uncertainty and validation layer, ONNX/browser engine, GeoTIFF export, web app.", {})]],
     size=8.6, color=INK, space_after=3)
notes(s6, "All references were consulted for the design; the baseline network is openly licensed (CC0-1.0).")

prs.save(str(OUT))
print("saved", OUT)
